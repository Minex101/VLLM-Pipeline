from openpyxl import load_workbook
from vllm import LLM, SamplingParams
from transformers import AutoProcessor
from PIL import Image
import csv

# Parameters
GT_PATH = "./dataset/dataset.xlsx"
IMAGE_PATH = "./dataset/"
OUTPUT_PATH = "./results.csv"

# -- MODELS -- 
MODEL = "Qwen/Qwen2.5-VL-72B-Instruct-AWQ" # Use the Qwen_sif img
# MODEL = "OpenGVLab/InternVL2_5-26B" # Use the Qwen_sif img

# -- PROMPTS -- 
PROMPT_TEXT = ( # Zero Shot Prompt (Baseline)
    "Look at this product image. Estimate the object's weight in grams and its "
    "largest dimension in centimeters. Give each as a min-max range.\n\n"
    "Respond with ONLY numbers separated by commas, in this exact order: "
    "weight_min_g,weight_max_g,largest_dim_min_cm,largest_dim_max_cm"
)

class EstimatorPipeline:
    def __init__(self, model_name):

        self.model_name = model_name
        self.llm = LLM(
            model=model_name,
            tensor_parallel_size=2,
            max_model_len=8192,
            gpu_memory_utilization=0.85,
            limit_mm_per_prompt={"image": 1, "video": 0},
            enforce_eager=True,
            # trust_remote_code=True # True; for InternVL Model
        )

        self.processor = AutoProcessor.from_pretrained(
            model_name,
            # trust_remote_code=True # True; for InternVL Model
        )
        self.sampling_params = SamplingParams(temperature=0, max_tokens=80)
    
    def build_prompt(self): # Enable for Qwen Model
        messages = [
            {
                "role": "user",
                "content": [
                    {"type": "image"},
                    {"type": "text", "text": PROMPT_TEXT},
                ],
            }
        ]
        return self.processor.apply_chat_template(
            messages, tokenize=False, add_generation_prompt=True
        )

    # def build_prompt(self):  # Enable for InternVL Model
    #     messages = [
    #         {
    #             "role": "user",
    #             "content": "<image>\n" + PROMPT_TEXT
    #         }
    #     ]

    #     return self.processor.apply_chat_template(
    #         messages,
    #         tokenize=False,
    #         add_generation_prompt=True
    #     )
    
    def estimate_batch(self, image_paths):
        inputs = []
        for path in image_paths:
            image = Image.open(path).convert("RGB")
            prompt = self.build_prompt()
            inputs.append({"prompt": prompt, "multi_modal_data": {"image": image}})

        outputs = self.llm.generate(inputs, self.sampling_params)
        return [o.outputs[0].text.strip() for o in outputs]
    
def get_image_name(ws):
    image_names = []
    for row in ws.iter_rows(min_col=6, max_col=6, min_row=4, max_row=5, values_only=True):
        image_names.append(row[0])
    return image_names

 # ---

def main():
    wb = load_workbook(GT_PATH)
    ws = wb["Objects"]

    pipeline = EstimatorPipeline(MODEL)

    # Get the image names to open the image and predict
    image_names = get_image_name(ws)

    image_paths = [IMAGE_PATH + name for name in image_names]
    results = pipeline.estimate_batch(image_paths)

    combined_results = []
    for name, result in zip(image_names, results):
        w_min, w_max, d_min, d_max = result.split(",")
        combined_results.append({
            "image_name": name,
            "weight_min_g": w_min.strip(),
            "weight_max_g": w_max.strip(),
            "largest_dim_min_cm": d_min.strip(),
            "largest_dim_max_cm": d_max.strip(),
        })

    with open(OUTPUT_PATH, "w", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "image_name",
                "weight_min_g",
                "weight_max_g",
                "largest_dim_min_cm",
                "largest_dim_max_cm",
            ],
        )
        writer.writeheader()
        writer.writerows(combined_results)

    print(f"\nSaved {len(combined_results)} results to {OUTPUT_PATH}")

if __name__ == "__main__":
    main()